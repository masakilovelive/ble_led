import argparse
import hashlib
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


MATH_RE = re.compile(
    r"(\$\$.*?\$\$|\\\[.*?\\\]|\\\(.*?\\\)|\$(?!\$).*?\$)",
    re.DOTALL,
)


def _split_math(text):
    parts = []
    last = 0
    for match in MATH_RE.finditer(text):
        if match.start() > last:
            parts.append(("text", text[last:match.start()]))
        parts.append(("math", match.group(0)))
        last = match.end()
    if last < len(text):
        parts.append(("text", text[last:]))
    return parts


def _is_display(token):
    return token.startswith("$$") or token.startswith("\\[")


def _strip_delims(token):
    if token.startswith("$$") and token.endswith("$$"):
        return token[2:-2]
    if token.startswith("\\[") and token.endswith("\\]"):
        return token[2:-2]
    if token.startswith("\\(") and token.endswith("\\)"):
        return token[2:-2]
    if token.startswith("$") and token.endswith("$"):
        return token[1:-1]
    return token


def _resolve_tex_bin(tex_bin_arg):
    candidates = []
    if tex_bin_arg:
        candidates.append(Path(tex_bin_arg))
    env_bin = os.environ.get("TEXLIVE_BIN")
    if env_bin:
        candidates.append(Path(env_bin))
    candidates.append(Path(r"C:\texlive\2025\bin\windows"))
    for candidate in candidates:
        if (candidate / "latex.exe").exists() and (candidate / "dvipng.exe").exists():
            return candidate
    return None


def resolve_tex_bin(tex_bin_arg):
    return _resolve_tex_bin(tex_bin_arg)


def _run_checked(cmd, cwd):
    proc = subprocess.run(
        cmd,
        cwd=cwd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        check=False,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if proc.returncode != 0:
        raise RuntimeError(
            "Command failed:\n{}\n\n{}".format(" ".join(cmd), proc.stdout)
        )


def _render_math(tex, is_display, out_dir, tex_bin, dpi):
    tex = tex.strip()
    key = "{}|{}|{}".format("display" if is_display else "inline", dpi, tex)
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:16]
    out_path = Path(out_dir) / ("math_{}.png".format(digest))
    if out_path.exists():
        return out_path

    tmpdir = Path(tempfile.mkdtemp(prefix="texmath_"))
    try:
        tex_path = tmpdir / "math.tex"
        tex_src = (
            "\\documentclass{article}\n"
            "\\usepackage{amsmath,amssymb,bm}\n"
            "\\pagestyle{empty}\n"
            "\\begin{document}\n"
        )
        if is_display:
            tex_src += "\\[\n{}\n\\]\n".format(tex)
        else:
            tex_src += "${}$\n".format(tex)
        tex_src += "\\end{document}\n"
        tex_path.write_text(tex_src, encoding="utf-8")

        latex_cmd = [str(tex_bin / "latex.exe")] if tex_bin else ["latex"]
        dvipng_cmd = [str(tex_bin / "dvipng.exe")] if tex_bin else ["dvipng"]

        _run_checked(
            latex_cmd
            + [
                "-interaction=nonstopmode",
                "-halt-on-error",
                "-output-directory",
                str(tmpdir),
                str(tex_path),
            ],
            cwd=tmpdir,
        )

        dvi_path = tmpdir / "math.dvi"
        _run_checked(
            dvipng_cmd
            + [
                "-T",
                "tight",
                "-D",
                str(dpi),
                "-bg",
                "Transparent",
                "-o",
                str(out_path),
                str(dvi_path),
            ],
            cwd=tmpdir,
        )
        return out_path
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


def _safe_add_paragraph(doc, style_name):
    if not style_name:
        return doc.add_paragraph()
    try:
        return doc.add_paragraph(style=style_name)
    except Exception:
        return doc.add_paragraph()


def _paragraph_font_pt(paragraph, document, default_pt=11.0):
    size = None
    for run in paragraph.runs:
        if run.font is not None and run.font.size is not None:
            size = run.font.size
            break
    if size is None and paragraph.style is not None:
        try:
            size = paragraph.style.font.size
        except Exception:
            size = None
    if size is None:
        try:
            size = document.styles["Normal"].font.size
        except Exception:
            size = None
    if size is None:
        return default_pt
    try:
        return size.pt
    except Exception:
        return default_pt


def _scale_shape_to_height(shape, height_pt):
    try:
        ratio = None
        if shape.height:
            ratio = float(shape.width) / float(shape.height)
        target_height = Pt(height_pt)
        shape.height = target_height
        if ratio:
            shape.width = int(shape.height * ratio)
    except Exception:
        return


def _has_numbering(paragraph):
    try:
        ppr = paragraph._p.pPr
        return ppr is not None and ppr.numPr is not None
    except Exception:
        return False


def iter_docx_files(input_dir, recursive, suffix):
    pattern = "**/*.docx" if recursive else "*.docx"
    for path in input_dir.glob(pattern):
        if path.name.startswith("~$"):
            continue
        if suffix and path.stem.endswith(suffix):
            continue
        yield path


def build_output_path(input_path, input_root, output_dir, suffix):
    stem = input_path.stem + (suffix or "")
    if output_dir is None:
        return input_path.with_name(stem + input_path.suffix)
    rel = input_path.relative_to(input_root)
    out_path = output_dir / rel
    return out_path.with_name(stem + input_path.suffix)


def convert_docx_batch(
    input_dir,
    output_dir,
    tex_bin,
    dpi_inline,
    dpi_display,
    inline_scale,
    display_scale,
    cache_dir,
    recursive,
    suffix,
    skip_existing,
):
    input_dir = Path(input_dir)
    output_dir = Path(output_dir) if output_dir else None
    cache_dir = Path(cache_dir) if cache_dir else input_dir / ".math_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)

    results = []
    errors = []
    for docx_path in iter_docx_files(input_dir, recursive, suffix):
        out_path = build_output_path(docx_path, input_dir, output_dir, suffix)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        if skip_existing and out_path.exists():
            results.append((docx_path, out_path, "skipped"))
            continue
        try:
            convert_docx(
                docx_path,
                out_path,
                tex_bin,
                dpi_inline,
                dpi_display,
                inline_scale,
                display_scale,
                cache_dir,
            )
            results.append((docx_path, out_path, "ok"))
        except Exception as exc:
            errors.append((docx_path, str(exc)))
            results.append((docx_path, out_path, "error"))
    return results, errors


def convert_docx(
    input_path,
    output_path,
    tex_bin,
    dpi_inline,
    dpi_display,
    inline_scale,
    display_scale,
    cache_dir,
):
    in_doc = Document(str(input_path))
    out_doc = Document()
    cache_dir.mkdir(parents=True, exist_ok=True)

    for paragraph in in_doc.paragraphs:
        text = paragraph.text
        style_name = paragraph.style.name if paragraph.style is not None else None
        if _has_numbering(paragraph):
            style_name = "List Number"
        parts = _split_math(text)
        font_pt = _paragraph_font_pt(paragraph, in_doc)

        current_par = None
        for kind, chunk in parts:
            if kind == "text":
                if chunk:
                    if current_par is None:
                        current_par = _safe_add_paragraph(out_doc, style_name)
                    current_par.add_run(chunk)
                continue

            is_display = _is_display(chunk)
            math_body = _strip_delims(chunk)
            dpi = dpi_display if is_display else dpi_inline
            img_path = _render_math(
                math_body, is_display, cache_dir, tex_bin, dpi
            )

            if is_display:
                if current_par is not None and not current_par.text:
                    current_par = None
                display_par = _safe_add_paragraph(out_doc, style_name)
                display_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                display_run = display_par.add_run()
                shape = display_run.add_picture(str(img_path))
                _scale_shape_to_height(shape, font_pt * display_scale)
                current_par = None
            else:
                if current_par is None:
                    current_par = _safe_add_paragraph(out_doc, style_name)
                inline_run = current_par.add_run()
                shape = inline_run.add_picture(str(img_path))
                _scale_shape_to_height(shape, font_pt * inline_scale)

        if current_par is None and not parts:
            _safe_add_paragraph(out_doc, style_name)

    out_doc.save(str(output_path))


def main():
    parser = argparse.ArgumentParser(
        description="Convert LaTeX fragments in a DOCX into rendered images."
    )
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", nargs="?", type=Path)
    parser.add_argument("--tex-bin", type=str, default=None)
    parser.add_argument("--dpi-inline", type=int, default=200)
    parser.add_argument("--dpi-display", type=int, default=250)
    parser.add_argument("--cache-dir", type=Path, default=None)
    parser.add_argument("--inline-scale", type=float, default=1.0)
    parser.add_argument("--display-scale", type=float, default=1.0)
    parser.add_argument("--recursive", action="store_true")
    parser.add_argument("--output-dir", type=Path, default=None)
    parser.add_argument("--suffix", type=str, default="_compiled")
    parser.add_argument("--skip-existing", action="store_true")
    args = parser.parse_args()

    tex_bin = _resolve_tex_bin(args.tex_bin)
    cache_dir = (
        args.cache_dir
        if args.cache_dir is not None
        else args.input_docx.parent / ".math_cache"
    )
    if args.input_docx.is_dir():
        output_dir = args.output_dir
        if output_dir is not None:
            output_dir.mkdir(parents=True, exist_ok=True)
        results, errors = convert_docx_batch(
            args.input_docx,
            output_dir,
            tex_bin,
            args.dpi_inline,
            args.dpi_display,
            args.inline_scale,
            args.display_scale,
            cache_dir,
            args.recursive,
            args.suffix,
            args.skip_existing,
        )
        for src, out, status in results:
            print("{}: {} -> {}".format(status, src, out))
        if errors:
            raise SystemExit("Errors: {}".format(len(errors)))
        return

    if args.output_docx is None:
        args.output_docx = args.input_docx.with_name(
            args.input_docx.stem + args.suffix + args.input_docx.suffix
        )
    convert_docx(
        args.input_docx,
        args.output_docx,
        tex_bin,
        args.dpi_inline,
        args.dpi_display,
        args.inline_scale,
        args.display_scale,
        cache_dir,
    )


if __name__ == "__main__":
    main()
